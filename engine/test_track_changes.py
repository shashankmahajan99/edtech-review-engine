"""
Tests for native OpenXML Track Changes generation.
"""

import docx
from pathlib import Path
from engine.openxml_track_changes import WordTrackChangesEngine

def test_track_changes():
    test_dir = Path(__file__).resolve().parent.parent / "output"
    test_dir.mkdir(exist_ok=True)
    raw_doc_path = test_dir / "test_raw_manuscript.docx"
    tracked_doc_path = test_dir / "test_tracked_manuscript.docx"

    # Step 1: Create a realistic Word document
    doc = docx.Document()
    doc.add_heading("Admissions Personal Statement: Candidate 409", level=1)
    doc.add_paragraph("Ever since I was a child, I have always been deeply passionate about computer science.")
    doc.add_paragraph("During my sophomore year, I worked on a machine learning project that predicted housing prices using regression models.")
    doc.add_paragraph("This university is the ideal place for me because of its world-class faculty and facilities.")
    doc.save(str(raw_doc_path))

    # Step 2: Define line edits
    edits = [
        {
            "original_text": "I have always been deeply passionate about",
            "suggested_replacement": "my academic curiosity crystallized around",
            "explanation": "Eliminates generic childhood passion cliché in favor of active intellectual inquiry."
        },
        {
            "original_text": "This university is the ideal place for me because of its world-class faculty and facilities",
            "suggested_replacement": "The departmental emphasis on distributed architectures under Dr. [Faculty] directly aligns with my research objectives",
            "explanation": "Substitutes generic institutional praise with precise faculty alignment."
        }
    ]

    # Step 3: Run OpenXML track changes engine
    engine = WordTrackChangesEngine(author="Apex Admissions Editorial Desk")
    engine.process_document(str(raw_doc_path), str(tracked_doc_path), edits)

    # Step 4: Verify generated XML
    doc_tracked = docx.Document(str(tracked_doc_path))
    full_xml = "".join([p._p.xml for p in doc_tracked.paragraphs])

    has_ins = "<w:ins" in full_xml
    has_del = "<w:del" in full_xml
    has_author = "Apex Admissions Editorial Desk" in full_xml

    print("=" * 60)
    print("  OPENXML TRACK CHANGES VERIFICATION")
    print("=" * 60)
    print(f"Generated Docx Path : {tracked_doc_path}")
    print(f"Has <w:ins> tag     : {has_ins}")
    print(f"Has <w:del> tag     : {has_del}")
    print(f"Has Author Name     : {has_author}")
    print("=" * 60)

    assert has_ins and has_del and has_author, "Failed: Native OpenXML elements missing!"
    print("[PASS] True Word Track Changes XML successfully generated and verified!")

if __name__ == "__main__":
    test_track_changes()
