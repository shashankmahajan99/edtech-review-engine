"""Incorrect diagnoses and copied closing-note prose exposed by a scan export."""
from pathlib import Path
import tempfile

import docx
from lxml import etree

from engine.review.analysis import AnalysisResult, CONCLUSION, Finding, THESIS, analyze, quoted_passages
from engine.review.docmodel import DocumentModel, Paragraph, load_document, reference_heading_span, w
from engine.review.endnote import compose_endnote
from engine.review.persona import CommentComposer
from engine.review.selection import SelectedComment, select_comments


def draft(directory, paragraphs):
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    path = Path(directory) / 'draft.docx'
    document.save(path)
    return path


def test_joined_source_heading_stops_rhetorical_analysis():
    with tempfile.TemporaryDirectory() as directory:
        path = draft(directory, [
            'This discussion explains how students learn through examples and practice. ' * 5,
            'Students learn by comparing examples and explaining what they show. ' * 5,
            'These examples show why learning requires practice and explanation. ' * 5,
            'Works Cited Smith, Avery. “A Study.” Example Journal, 2025, https://example.test/study',
            'https://example.test/another-source'])
        model = load_document(str(path), academic_style='mla')
        assert model.paragraphs[3].region == 'references'
        assert model.paragraphs[4].region == 'references'
        result = analyze(model)
        assert any(f.rule_id == 'reference_heading_joined' for f in result.findings)
        assert any(f.category == CONCLUSION for f in result.findings)
        assert all(f.paragraph.region == 'body' for f in result.findings if f.category == CONCLUSION)


def test_reference_word_in_ordinary_prose_is_not_a_heading():
    assert reference_heading_span('References to the earlier discussion help connect these points.') is None
    assert reference_heading_span('Works Cited') == (0, 11)


def test_short_titles_are_consumed_before_finding_long_quotations():
    text = 'Although "Morning" by One Band and "Evening" by Another Band differ, both concern change.'
    assert list(quoted_passages(text)) == []
    curly = 'Although “Morning” by One Band and “Evening” by Another Band differ, both concern change.'
    assert list(quoted_passages(curly)) == []


def test_actual_quote_is_not_confused_with_preceding_short_title():
    quote = '“The morning light returns and gives the traveler another chance to begin”'
    text = 'The song “Morning” includes these words: ' + quote + '.'
    assert [m.group() for m in quoted_passages(text)] == [quote]


def quote_draft(directory, reference, locator=''):
    return draft(directory, [
        'The discussion concerns how evidence supports an argument in this paper. ' * 5,
        'Smith writes, “The findings give readers another reason to consider how the explanation works”' +
        locator + '. The passage introduces a point for the discussion. ' * 3,
        'The evidence and explanation belong together in the discussion. ' * 5,
        'Works Cited', reference])


def test_web_quotation_does_not_require_invented_page_numbers():
    with tempfile.TemporaryDirectory() as directory:
        path = quote_draft(directory, 'Smith, Avery. “A Web Article.” Example Site, 2025, https://example.test/article')
        rules = {f.rule_id for f in analyze(load_document(str(path), academic_style='mla')).findings}
        assert 'quote_no_page' not in rules


def test_known_paginated_source_can_receive_a_locator_note():
    with tempfile.TemporaryDirectory() as directory:
        path = quote_draft(directory, 'Smith, Avery. “A Study.” Example Journal, 2025, pp. 10–20.')
        model = load_document(str(path), academic_style='mla')
        finding = next(f for f in analyze(model).findings if f.rule_id == 'quote_no_page')
        assert finding.facts['source'] == 'Smith'
        body = CommentComposer(model).compose(SelectedComment(finding, finding.category, True))
        assert 'Smith' in body
        assert 'edition you used' in body


def test_number_only_locator_is_recognized():
    with tempfile.TemporaryDirectory() as directory:
        path = quote_draft(directory, 'Smith, Avery. “A Study.” Example Journal, 2025, pp. 10–20.', ' (14)')
        result = analyze(load_document(str(path), academic_style='mla'))
        assert not any(f.rule_id == 'quote_no_page' for f in result.findings)


def test_paginated_source_elsewhere_does_not_relabel_a_web_quotation():
    with tempfile.TemporaryDirectory() as directory:
        path = draft(directory, [
            'This discussion considers how several sources support the argument. ' * 5,
            'Smith discusses an earlier study. Jones writes, “The findings give readers another reason '
            'to consider how the explanation works”. The comparison helps readers follow the discussion. ' * 3,
            'The conclusion connects the different sources with the argument. ' * 5,
            'Works Cited',
            'Smith, Avery. “A Study.” Example Journal, 2025, pp. 10–20.',
            'Jones, Avery. “A Web Article.” Example Site, 2025, https://example.test/article'])
        result = analyze(load_document(str(path), academic_style='mla'))
        assert not any(f.rule_id == 'quote_no_page' for f in result.findings)


def test_comparison_topic_feedback_respects_an_explicit_analytical_feature():
    with tempfile.TemporaryDirectory() as directory:
        for opening, expected in [
            ('“Morning” and “Evening” both use another example of independence by showing how their lives are good.', True),
            ('“Morning” and “Evening” both use another example of independence through contrasting tone and imagery.', False),
        ]:
            path = draft(directory, [
                'This essay compares two songs and the different choices their speakers describe. ' * 5,
                opening + ' The examples in the discussion show the choices made by each speaker. ' * 5,
                'The conclusion connects the speakers and their choices in the songs. ' * 5])
            model = load_document(str(path), academic_style='mla')
            findings = [f for f in analyze(model).findings if f.rule_id == 'comparison_topic_focus']
            assert bool(findings) == expected
            if findings:
                body = CommentComposer(model).compose(SelectedComment(findings[0], findings[0].category, True))
                assert '“Morning” and “Evening”' in body
                assert 'be more specific' not in body


def test_analytical_topic_sentence_does_not_need_a_modal_verb():
    with tempfile.TemporaryDirectory() as directory:
        opening = 'Both poems demonstrate independence through the speakers’ different choices.'
        path = draft(directory, [
            'This discussion compares two poems about independence and personal choices. ' * 5,
            opening + ' The first speaker refuses an invitation while the second accepts help. ' * 5,
            'The speakers reach different conclusions about accepting help and making choices. ' * 5])
        result = analyze(load_document(str(path), academic_style='mla'))
        assert not any(f.rule_id == 'para_no_claim' and f.paragraph.index == 1 for f in result.findings)


def test_word_overlap_alone_cannot_diagnose_paragraph_drift():
    with tempfile.TemporaryDirectory() as directory:
        path = draft(directory, [
            'The opening introduces a comparison of education and participation. ' * 5,
            'Students consider the cost of attendance. A family may need financial support. '
            'Working fewer hours can leave time for studying. Public funding offers one possible response. '
            'These examples concern the resources available to a learner.',
            'The conclusion returns to the connection between education and participation. ' * 5])
        assert not any(f.rule_id == 'para_drift' for f in analyze(load_document(str(path))).findings)


def test_thesis_precision_note_does_not_claim_the_argument_is_absent():
    text = 'The two songs present independence through priorities, good lives and escaping unpleasant things.'
    paragraph = Paragraph(0, etree.Element(w('p')), text, [], 'Normal')
    model = DocumentModel('test.docx', [paragraph], None, 'test')
    finding = Finding('thesis_vague', THESIS, paragraph, 0, len(text), {'vague': ['things']}, 82)
    item = SelectedComment(finding, THESIS, True)
    body = CommentComposer(model).compose(item)
    assert '“things”' in body
    assert 'not clearly established' not in body
    assert 'already stated' in body


def test_closing_plan_expands_the_revision_without_pasting_the_comment():
    text = 'The two songs present independence through priorities, good lives and escaping unpleasant things.'
    paragraph = Paragraph(0, etree.Element(w('p')), text, [], 'Normal')
    front = Paragraph(1, etree.Element(w('p')), 'Avery Smith Mrs. Jones ENG 111', [], 'Normal', region='front_matter')
    model = DocumentModel('test.docx', [front, paragraph], None, 'test')
    finding = Finding('thesis_vague', THESIS, paragraph, 0, len(text), {'vague': ['things']}, 82)
    item = SelectedComment(finding, THESIS, True)
    body = CommentComposer(model).compose(item)
    analysis = AnalysisResult([finding], 'academic', [], {}, [])
    note = compose_endnote(model, analysis, composed=[(item, body)])
    assert note.paragraphs[0] == 'Dear Avery,'
    assert body not in ' '.join(note.paragraphs)
    assert note.evidence[0]['quote'] == text
    assert 'Work through these areas in the order listed.' not in ' '.join(note.paragraphs)


def test_repeated_vague_word_is_addressed_once_per_review():
    with tempfile.TemporaryDirectory() as directory:
        path = draft(directory, [
            'The discussion introduces a comparison of the things that affect learning. ' * 5,
            'These things relate to the resources available to students. ' * 5,
            'Other things relate to the time needed to practice. ' * 5])
        model = load_document(str(path))
        selected = select_comments(model, analyze(model), comprehensive=True)
        assert sum(s.finding.rule_id == 'informal_diction' for s in selected) <= 1


def test_degrees_and_urls_are_not_missing_sentence_spaces():
    with tempfile.TemporaryDirectory() as directory:
        path = draft(directory, ['The author Avery Smith, Ed.D., discusses this research. '
                                 'The article appears at https://example.test/Research.Article. ' * 5])
        result = analyze(load_document(str(path)))
        assert not any(f.rule_id in {'missing_space_after_punct', 'typo_cluster'} for f in result.findings)


def test_pipeline_rejects_rhetorical_feedback_on_a_source_entry():
    from unittest.mock import patch
    import pytest
    from engine.review.pipeline import PipelineError, review_document
    with tempfile.TemporaryDirectory() as directory:
        path = draft(directory, ['This discussion explains the evidence used to develop an argument. ' * 8,
                                 'Works Cited', 'Smith, Avery. “A Study.” Example Journal, 2025, pp. 10–20.'])
        def invalid_analysis(model):
            paragraph = model.references[-1]
            finding = Finding('conclusion_summary', CONCLUSION, paragraph, 0, len(paragraph.text), {}, 76)
            return AnalysisResult([finding], 'academic', [], {}, [])
        with patch('engine.review.pipeline.analyze', side_effect=invalid_analysis):
            with pytest.raises(PipelineError, match='source entries'):
                review_document(path, Path(directory) / 'out')
        assert not list(Path(directory).rglob('*_commented.docx'))
