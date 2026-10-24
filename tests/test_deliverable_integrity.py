"""Deliverable integrity: preservation, in-place redlines, full coverage."""
from __future__ import annotations

import sys
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import docx  # noqa: E402

from engine.review.analysis import analyze  # noqa: E402
from engine.review.docmodel import load_document  # noqa: E402
from engine.review.selection import select_comments  # noqa: E402
from engine.review.trackchanges import apply_line_edits  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures"
OUT = ROOT / "tests" / "_out"


class TestRedlinePreservesFormatting(unittest.TestCase):
    def setUp(self):
        OUT.mkdir(parents=True, exist_ok=True)
        self.source = OUT / "redline_input.docx"
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("This is ")
        emphasised = paragraph.add_run("very important")
        emphasised.bold = True
        emphasised.italic = True
        paragraph.add_run(" text, and the reason for that is because it repeats.")
        document.add_paragraph("A second paragraph that must not move or change.")
        document.save(str(self.source))

    def test_only_document_xml_changes(self):
        output = OUT / "redline_output.docx"
        applied = apply_line_edits(
            str(self.source), str(output),
            [{"original_text": "the reason for that is because",
              "suggested_replacement": "because"}],
        )
        self.assertEqual(applied, 1)

        with zipfile.ZipFile(self.source) as before, zipfile.ZipFile(output) as after:
            self.assertEqual(set(before.namelist()), set(after.namelist()))
            changed = [n for n in before.namelist()
                       if n != "word/document.xml" and before.read(n) != after.read(n)]
            self.assertEqual(changed, [])

        result = docx.Document(str(output))
        runs = [run for run in result.paragraphs[0].runs if run.text.strip()]
        bold = [run for run in runs if run.bold and run.italic]
        self.assertTrue(bold, "bold/italic run was lost")

        model = load_document(str(output))
        accepted = model.paragraphs[0].text
        self.assertIn("because it repeats", accepted)
        self.assertNotIn("the reason for that is because", accepted)
        self.assertEqual(model.paragraphs[1].text,
                         "A second paragraph that must not move or change.")

    def test_tracked_changes_are_native(self):
        output = OUT / "redline_native.docx"
        apply_line_edits(str(self.source), str(output),
                         [{"original_text": "very important",
                           "suggested_replacement": "central"}])
        with zipfile.ZipFile(output) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("<w:ins ", xml)
        self.assertIn("<w:del ", xml)
        self.assertIn("central", xml)


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestStudioPreservesOriginal(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        from portal.server import app

        cls.client = TestClient(app)
        cls.fixture = FIXTURES / "andrea_raw.docx"
        response = cls.client.post(
            "/api/thinkingstorm/execute",
            data={"student_name": "Andrea Goblet", "reviewer_name": "Raunak Banga",
                  "reviewer_initials": "RB", "academic_style": "mla", "text": "",
                  "custom_comments_text": "", "custom_endnote_text": ""},
            files={"file": (cls.fixture.name, cls.fixture.read_bytes())},
        )
        cls.status = response.status_code
        cls.payload = response.json() if response.status_code == 200 else {}

    def test_review_succeeds(self):
        self.assertEqual(self.status, 200, self.payload)
        self.assertGreaterEqual(len(self.payload.get("comments", [])), 5)

    def test_redline_keeps_every_part(self):
        job = self.payload["job_id"]
        annotated = self.client.get(f"/download/docx/{job}")
        self.assertEqual(annotated.status_code, 200)
        path = OUT / f"{job}_annotated.docx"
        path.write_bytes(annotated.content)

        with zipfile.ZipFile(self.fixture) as before, zipfile.ZipFile(path) as after:
            self.assertEqual(set(before.namelist()), set(after.namelist()))
            changed = [n for n in before.namelist()
                       if n != "word/document.xml" and before.read(n) != after.read(n)]
            self.assertEqual(changed, [])
            xml = after.read("word/document.xml").decode("utf-8")
        self.assertIn("<w:ins ", xml)
        self.assertIn("<w:del ", xml)

    def test_commented_keeps_original_text(self):
        from engine.review.docmodel import build_paragraph

        job = self.payload["job_id"]
        commented = self.client.get(f"/download/commented/{job}")
        self.assertEqual(commented.status_code, 200)
        path = OUT / f"{job}_commented.docx"
        path.write_bytes(commented.content)

        with zipfile.ZipFile(self.fixture) as before, zipfile.ZipFile(path) as after:
            self.assertTrue(set(before.namelist()).issubset(set(after.namelist())))

        original = load_document(str(self.fixture))
        produced = load_document(str(path))
        for index, paragraph in enumerate(original.paragraphs):
            self.assertEqual(paragraph.text, produced.paragraphs[index].text,
                             f"paragraph {index} changed")


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestFullCoverage(unittest.TestCase):
    def test_every_block_with_issues_is_annotated(self):
        for name in ("janeen_raw.docx", "kiran_raw.docx"):
            model = load_document(str(FIXTURES / name))
            result = analyze(model)
            selected = select_comments(model, result, comprehensive=True)
            covered_blocks = 0
            for block in result.blocks:
                ids = {p.index for p in block.paragraphs}
                if any(f.paragraph.index in ids for f in result.findings) and \
                        any(s.finding.paragraph.index in ids for s in selected):
                    covered_blocks += 1
            blocks_with_findings = sum(
                1 for block in result.blocks
                if any(f.paragraph.index in {p.index for p in block.paragraphs}
                       for f in result.findings)
            )
            self.assertEqual(covered_blocks, blocks_with_findings, name)
            self.assertGreaterEqual(len(selected), 8, name)


if __name__ == "__main__":
    unittest.main(verbosity=2)
