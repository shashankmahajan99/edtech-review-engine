"""Content regressions: complete advice, diagnostic fidelity and provenance."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from engine.review.analysis import AnalysisResult, Finding, THESIS, CONCLUSION, PUNCT, analyze
from engine.review.docmodel import DocumentModel, Paragraph, classify_regions, detect_style, w
from engine.review.endnote import compose_endnote
from engine.review.persona import CommentComposer, CompositionError, TEMPLATES
from engine.review.selection import SelectedComment, _premerge
from engine.review import kb


def model_and_item(rule="thesis_vague", category=THESIS, text="Schools should change their schedules.", facts=None):
    para = Paragraph(0, etree.Element(w("p")), text, [], "Normal")
    model = DocumentModel("test.docx", [para], None, "test-fingerprint")
    finding = Finding(rule, category, para, 0, len(text), facts or {}, 80)
    return model, SelectedComment(finding, category, True)


class TestCommentContent(unittest.TestCase):
    def test_advice_is_never_dropped_by_variant_hash(self):
        model, item = model_and_item()
        with patch.dict(TEMPLATES, {"thesis_vague": [("The claim needs a reason.", "Add the reason supporting your position.")]}), \
                patch.object(kb, "patterns_for_rule", return_value=([], [])):
            for digest in range(60):
                composer = CommentComposer(model)
                with patch.object(composer, "_variant_key", return_value=f"{digest:x}"):
                    body = composer.compose(item)
                self.assertIn("Add the reason supporting your position.", body)

    def test_exhausted_pool_keeps_rendered_facts_and_advice(self):
        model, item = model_and_item("wordiness", "SENTENCE CLARITY", facts={"phrase": "in order to"})
        with patch.dict(TEMPLATES, {"wordiness": [("Replace “{phrase}”.", "Use a shorter phrase.")]}):
            composer = CommentComposer(model)
            bodies = [composer.compose(item) for _ in range(5)]
        self.assertEqual(len(set(bodies)), 1)
        self.assertIn("in order to", bodies[-1])
        self.assertNotIn("{phrase}", bodies[-1])
        self.assertIn("Use a shorter phrase.", bodies[-1])

    def test_missing_required_fact_fails_instead_of_inventing_a_name(self):
        model, item = model_and_item("org_name_mismatch", "APA")
        with patch.dict(TEMPLATES, {"org_name_mismatch": [("Match {long} to {short}.", "Check both names.")]}):
            with self.assertRaises(CompositionError):
                CommentComposer(model).compose(item)

    def test_unknown_rule_has_no_generic_fallback(self):
        model, item = model_and_item("unknown_diagnostic")
        with self.assertRaises(CompositionError):
            CommentComposer(model).compose(item)

    def test_student_braces_are_preserved_without_second_format_pass(self):
        model, item = model_and_item()
        with patch.object(kb, "patterns_for_rule", return_value=([
            ("Your notation “{x}” needs an explanation.", "Define the symbol for the reader.")], [])), \
                patch.dict(TEMPLATES, {"thesis_vague": [("{missing}", "{missing}")]}):
            body = CommentComposer(model).compose(item)
        self.assertIn("{x}", body)

    def test_category_match_does_not_authorize_a_different_diagnosis(self):
        model, item = model_and_item("conclusion_summary", CONCLUSION)
        unbound = {"moves": [{"id": "conclusion", "patterns": [{
            "observation": "A new point appears at the end.", "advice": "Move it into the body."
        }]}]}
        with patch.object(kb, "load", return_value=unbound):
            self.assertEqual(kb.patterns_for_rule(item.finding.rule_id, model, item.finding), ([], []))

    def test_knowledge_base_rule_binding_is_respected(self):
        model, item = model_and_item("conclusion_summary", CONCLUSION)
        data = {"moves": [{"id": "conclusion", "patterns": [{
            "rules": ["conclusion_summary"], "observation": "The ending repeats the discussion.",
            "advice": "Explain the implication."
        }]}]}
        with patch.object(kb, "load", return_value=data):
            self.assertTrue(kb.patterns_for_rule(item.finding.rule_id, model, item.finding)[1])


class TestEndnoteEvidence(unittest.TestCase):
    def setUp(self):
        self.model, self.item = model_and_item("double_period", PUNCT, "The lesson ends here..")
        self.analysis = AnalysisResult([self.item.finding], "academic", [], {}, [])
        self.composed = [(self.item, "There is a doubled period here. Reduce it to a single mark.")]

    def test_only_delivered_findings_supply_revision_areas(self):
        omitted = Finding("thesis_missing", THESIS, self.item.finding.paragraph, 0, 8, {}, 100)
        self.analysis.findings.append(omitted)
        note = compose_endnote(self.model, self.analysis, composed=self.composed)
        self.assertEqual(note.domains, ["Grammar and punctuation"])
        self.assertEqual(note.evidence[0]["rule"], "double_period")
        text = " ".join(note.paragraphs).lower()
        self.assertNotIn("thesis", text)
        self.assertNotIn("sources", text)
        self.assertNotIn("logical", text)
        self.assertIn("The lesson ends here..", " ".join(note.paragraphs))

    def test_explicit_student_name_overrides_front_matter(self):
        front = Paragraph(1, etree.Element(w("p")), "Wrong Person", [], "Normal", region="front_matter")
        self.model.paragraphs.insert(0, front)
        note = compose_endnote(self.model, self.analysis, student_name="Correct Student", composed=self.composed)
        self.assertEqual(note.paragraphs[0], "Dear Correct,")

    def test_selection_does_not_mutate_analysis_when_merging(self):
        para = self.item.finding.paragraph
        first = Finding("conclusion_summary", CONCLUSION, para, 0, len(para.text), {}, 90)
        second = Finding("double_period", PUNCT, para, len(para.text) - 3, len(para.text), {}, 60)
        merged = _premerge([first, second])
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0].rule_id, "conclusion_and_punctuation")
        self.assertEqual(first.rule_id, "conclusion_summary")
        self.assertEqual(first.category, CONCLUSION)
        self.assertEqual(second.rule_id, "double_period")


class TestDocumentRegions(unittest.TestCase):
    def test_later_heading_cannot_hide_earlier_student_prose(self):
        first = Paragraph(0, etree.Element(w("p")), "Taylor Caruso", [], "Normal")
        prose = Paragraph(1, etree.Element(w("p")), "A student explains the subject in this passage. " * 5, [], "Normal")
        late = Paragraph(2, etree.Element(w("p")), "Later Section", [], "Heading 1", is_heading=True)
        classify_regions([first, prose, late])
        self.assertEqual(prose.region, "body")
        self.assertNotEqual(late.region, "title")

    def test_singular_reference_heading_starts_reference_section(self):
        prose = Paragraph(0, etree.Element(w("p")), "An introduction with enough text to count as body prose. " * 4, [], "Normal")
        heading = Paragraph(1, etree.Element(w("p")), "Reference", [], "Normal")
        source = Paragraph(2, etree.Element(w("p")), "Smith, A. (2025). Title. https://example.test/source", [], "Normal")
        classify_regions([prose, heading, source])
        self.assertEqual(source.region, "references")

    def test_narrative_author_year_citation_supports_apa_detection(self):
        self.assertEqual(detect_style("Morrin et al. (2026) reported the results.", []), "apa")

    def test_clinical_report_does_not_receive_essay_or_title_page_diagnoses(self):
        model, item = model_and_item(text="An observation of the child's participation in daily activities. " * 12)
        title = Paragraph(1, etree.Element(w("p")), "Occupational Therapy Evaluation", [], "Normal", region="title")
        patient_date = Paragraph(2, etree.Element(w("p")), "DOB: 07/17/2025", [], "Normal", region="front_matter")
        model.paragraphs = [title, patient_date, item.finding.paragraph]
        model.style = "apa"
        result = analyze(model)
        self.assertEqual(result.essay_type, "report")
        rules = {f.rule_id for f in result.findings}
        self.assertFalse(rules & {"thesis_missing", "conclusion_summary", "apa_date_numeric", "apa_instructor_missing"})

    def test_annotated_bibliography_is_not_required_to_argue_between_entries(self):
        model, item = model_and_item(text="According to Smith (2025), the study reported changes in participation. " * 8)
        title = Paragraph(1, etree.Element(w("p")), "Annotated Bibliography", [], "Normal", region="title")
        model.paragraphs.insert(0, title)
        result = analyze(model)
        self.assertEqual(result.essay_type, "bibliography")
        rules = {f.rule_id for f in result.findings}
        self.assertFalse(rules & {"thesis_missing", "conclusion_summary", "topic_evidence_opener", "para_no_claim", "para_unsupported"})


class TestStudioQualityGate(unittest.TestCase):
    def test_validation_failure_does_not_call_legacy_generator(self):
        from fastapi import HTTPException
        from portal import server
        from engine.review.pipeline import PipelineError

        text = "Schools should review their schedules and explain the changes to students. " * 12
        with tempfile.TemporaryDirectory() as directory, \
                patch.object(server, "OUTPUT_DIR", Path(directory)), \
                patch("engine.review.pipeline.review_document", side_effect=PipelineError("test rejection")), \
                patch.object(server.ThinkingStormAutomator, "analyze_manuscript") as legacy:
            request = server.execute_thinkingstorm_pipeline(
                student_name="Taylor", reviewer_name="Reviewer", reviewer_initials="R",
                academic_style="mla", text=text, custom_comments_text="",
                custom_endnote_text="", file=None)
            import asyncio
            with self.assertRaises(HTTPException) as error:
                asyncio.run(request)
            self.assertEqual(error.exception.status_code, 422)
            self.assertIn("test rejection", error.exception.detail)
            legacy.assert_not_called()


class TestSparseFindings(unittest.TestCase):
    def test_one_verified_issue_is_delivered_without_filler(self):
        import docx
        from engine.review.pipeline import review_document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reflection.docx"
            document = docx.Document()
            document.add_paragraph("A student reflects on a particular classroom experience and describes the lesson learned. " * 7 + "The account ends here..")
            document.save(path)

            def one_finding(model):
                para = model.paragraphs[0]
                start = para.text.index("ends here..")
                finding = Finding("double_period", PUNCT, para, start, len(para.text), {}, 61)
                return AnalysisResult([finding], "reflection", [], {}, [])

            with patch("engine.review.pipeline.analyze", side_effect=one_finding):
                result = review_document(path, Path(directory) / "out")
            self.assertTrue(result.validation.ok, result.validation.errors)
            self.assertEqual(len(result.comments), 1)
            self.assertTrue(any("Few supported findings" in w for w in result.validation.warnings))
            self.assertEqual(sum(p.endswith(":") for p in result.endnote), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
