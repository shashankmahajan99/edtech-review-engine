"""Local inference must not weaken manuscript integrity or hide failures."""
import json
from pathlib import Path

import pytest
from docx import Document

from engine.review.analysis import analyze
from engine.review.docmodel import load_document
from engine.review.local_model import (
    LocalReviewError, apply_editorial_changes, chat, configuration, decode_comments, decode_endnote, examples_for, parse_model_json,
)
from engine.review.pipeline import PipelineError, review_document


@pytest.fixture
def manuscript(tmp_path):
    doc = Document()
    doc.add_paragraph('A comparison of local transport')
    doc.add_paragraph('Public buses reduce the number of cars on busy roads. '
                      'However, a service that runs only once an hour cannot reliably replace driving. '
                      'The city should compare the cost of running more buses with the savings from shorter traffic delays. '
                      'This proposal depends on the number of commuters who would switch to a more frequent service. '
                      'A trial on one busy route would help establish whether frequency changes travel choices. '
                      'The council could compare passenger numbers before and after the trial, while checking whether other routes lost passengers.')
    path = tmp_path / 'draft.docx'
    doc.save(path)
    return load_document(str(path))


def feedback(manuscript):
    para = next(p for p in manuscript.paragraphs if 'Public buses' in p.text)
    return {'comments': [{'paragraph': para.index,
                          'quote': 'The city should compare the cost of running more buses',
                          'category': 'ANALYSIS OF EVIDENCE',
                          'reason': 'The cost comparison is proposed without identifying what costs will be counted.',
                          'body': 'Which operating costs would this comparison include? Naming the costs of additional drivers and vehicles would make the proposal more precise and give the council a clearer basis for assessing the trial.'}]}


def test_local_anchor_must_match_exact_student_text(manuscript):
    proposed = feedback(manuscript)
    composed = decode_comments(proposed, manuscript)
    assert composed[0][0].finding.quote == proposed['comments'][0]['quote']
    proposed['comments'][0]['quote'] = 'This sentence was invented by the model.'
    with pytest.raises(LocalReviewError, match='anchor must occur exactly once'):
        decode_comments(proposed, manuscript)


def test_local_overlap_is_rejected(manuscript):
    proposed = feedback(manuscript)
    proposed['comments'].append(dict(proposed['comments'][0]))
    with pytest.raises(LocalReviewError, match='overlapping'):
        decode_comments(proposed, manuscript)


def test_missing_diagnosis_is_rejected(manuscript):
    proposed = feedback(manuscript)
    del proposed['comments'][0]['reason']
    with pytest.raises(LocalReviewError, match='missing evidence'):
        decode_comments(proposed, manuscript)


def test_specific_strength_does_not_require_an_invented_revision(manuscript):
    proposed = feedback(manuscript)
    proposed['comments'][0].update(kind='strength', reason='The draft identifies a cost comparison that would help the council evaluate the proposed trial.', body='Your cost comparison gives the council a concrete way to evaluate the proposed trial.')
    composed = decode_comments(proposed, manuscript)
    assert composed[0][0].finding.facts['kind'] == 'strength'


def test_three_specific_strengths_are_not_a_fatal_schema_error(manuscript):
    row = feedback(manuscript)['comments'][0]
    quotes = ['Public buses reduce the number of cars on busy roads.',
              'The city should compare the cost of running more buses',
              'A trial on one busy route would help establish whether frequency changes travel choices.']
    rows = [dict(row, quote=quote, kind='strength',
                 body='This choice gives the council a concrete way to assess the proposed trial.',
                 reason='The marked passage identifies a specific benefit or practical comparison in the proposal.') for quote in quotes]
    assert len(decode_comments({'comments': rows}, manuscript)) == 3


def test_redundant_heading_is_not_copied_as_feedback_prose(manuscript):
    proposed = feedback(manuscript)
    body = proposed['comments'][0]['body']
    proposed['comments'][0]['body'] = 'ANALYSIS OF EVIDENCE: ' + body
    assert decode_comments(proposed, manuscript)[0][1] == body


def test_final_editorial_rewrite_cannot_move_an_anchor(manuscript):
    rows = feedback(manuscript)['comments']
    revised = apply_editorial_changes(rows, {'changes': [{'id': 0, 'action': 'rewrite',
        'body': rows[0]['body'], 'reason': rows[0]['reason'], 'quote': 'An invented anchor.'}]})
    assert revised['comments'][0]['quote'] == rows[0]['quote']
    assert decode_comments(revised, manuscript)


def test_final_editorial_check_rejects_unknown_comment_id(manuscript):
    with pytest.raises(LocalReviewError, match='invalid or duplicate'):
        apply_editorial_changes(feedback(manuscript)['comments'], {'changes': [{'id': 50, 'action': 'remove'}]})


def test_checkpoint_from_another_manuscript_is_rejected(manuscript, tmp_path, monkeypatch):
    from engine.review.local_model import compose_local
    checkpoint = tmp_path / 'checkpoint.json'
    checkpoint.write_text(json.dumps({'fingerprint': 'different manuscript', 'model': 'test-model', 'review': feedback(manuscript)}))
    monkeypatch.setattr('engine.review.local_model.configuration', lambda: {'model': 'test-model'})
    monkeypatch.setattr('engine.review.local_model.examples_for', lambda *args: ([], 0))
    monkeypatch.setattr('engine.review.local_model.chat', lambda *args: pytest.fail('A mismatched checkpoint reached inference'))
    with pytest.raises(LocalReviewError, match='does not match'):
        compose_local(manuscript, analyze(manuscript), 'Reviewer', 'Student', checkpoint=checkpoint)


def test_example_retrieval_excludes_same_document_under_new_filename(manuscript, tmp_path):
    para = next(p for p in manuscript.paragraphs if 'Public buses' in p.text)
    rows = [dict(source='original-name.docx', manuscript_probe=para.text[:120],
                 category='ANALYSIS', quote='buses', feedback='Do not retrieve this original review.'),
            dict(source='other.docx', manuscript_probe='A completely different text.',
                 category='ANALYSIS', quote='buses', feedback='An unrelated voice example.')]
    index = tmp_path / 'examples.json'
    index.write_text(json.dumps(rows), encoding='utf8')
    retrieved, excluded = examples_for(manuscript, {'examples_path': str(index)})
    assert excluded == 1
    assert [r['feedback'] for r in retrieved] == ['An unrelated voice example.']


def test_cloud_endpoint_cannot_receive_local_manuscripts(tmp_path, monkeypatch):
    config = tmp_path / 'config.json'
    config.write_text(json.dumps({'provider': 'ollama', 'base_url': 'http://example.com'}))
    monkeypatch.setenv('TS_LOCAL_CONFIG', str(config))
    with pytest.raises(LocalReviewError, match='loopback'):
        configuration()


def test_local_failure_does_not_silently_use_old_templates(manuscript, tmp_path, monkeypatch):
    def unavailable(*args):
        raise LocalReviewError('Local runtime unavailable')
    monkeypatch.setattr('engine.review.local_model.compose_local', unavailable)
    with pytest.raises(PipelineError, match='runtime unavailable'):
        review_document(manuscript.path, tmp_path / 'output', backend='local')
    assert not list((tmp_path / 'output').glob('*.docx'))


def test_endnote_cannot_reference_nonexistent_feedback(manuscript):
    composed = decode_comments(feedback(manuscript), manuscript)
    note = {'opening': 'Revise the proposal.', 'areas': [
        {'heading': 'Evidence', 'comment_ids': [100], 'text': 'Invented criticism.'},
        {'heading': 'Focus', 'comment_ids': [0], 'text': 'Another priority.'}], 'closing': 'Thank you.'}
    with pytest.raises(LocalReviewError, match='approved comment'):
        decode_endnote(note, manuscript, composed, 'Reviewer', 'Student')


def test_formatter_keeps_one_greeting_in_a_model_letter(manuscript):
    composed = decode_comments(feedback(manuscript), manuscript)
    plan = ('Begin with a table that separates the proposed service changes from the outcomes you want to measure. '
            'Choose one route and state which observations would show whether the trial helped commuters. '
            'Explain who would collect those observations and how the council could compare them with the earlier service. '
            'Keep the initial test small enough that its results can inform a later decision.')
    note = {'opening': 'Hi Student, your next revision can make the proposal easier to assess.',
            'areas': [{'heading': 'Developing the trial', 'comment_ids': [0], 'text': plan},
                      {'heading': 'Explaining the comparison', 'comment_ids': [0],
                       'text': plan.replace('Begin with a table', 'Prepare a short outline').replace('Choose one route', 'Select a suitable route')}],
            'closing': 'Read the revised proposal once more for the relationship between its purpose and the proposed measurements.'}
    letter = decode_endnote(note, manuscript, composed, 'Reviewer', 'Student')
    assert letter.paragraphs[:2] == ['Dear Student,', 'Your next revision can make the proposal easier to assess.']
    assert 'Hi Student' not in ' '.join(letter.paragraphs)


@pytest.mark.parametrize('provider', ['ollama', 'strata', 'llamacpp'])
def test_local_transports_send_manuscript_to_expected_api_only(provider, monkeypatch):
    captured = {}
    content = json.dumps({'comments': []})
    response = ({'message': {'content': content}, 'done_reason': 'stop'} if provider == 'ollama'
                else {'choices': [{'message': {'content': content}, 'finish_reason': 'stop'}]})
    import io
    class Opener:
        def open(self, request, timeout):
            captured['url'] = request.full_url
            captured['payload'] = json.loads(request.data)
            return io.BytesIO(json.dumps(response).encode())
    monkeypatch.setattr('engine.review.local_model.build_opener', lambda *args: Opener())
    result = chat({'provider': provider, 'base_url': 'http://127.0.0.1:1234', 'model': 'local-test'},
                  [{'role': 'user', 'content': 'Manuscript text'}])
    assert result == {'comments': []}
    assert captured['url'] == 'http://127.0.0.1:1234' + ('/api/chat' if provider == 'ollama' else '/v1/chat/completions')
    assert captured['payload']['messages'][0]['content'] == 'Manuscript text'
    if provider == 'strata':
        assert captured['payload']['chat_template_kwargs']['enable_thinking'] is True
        assert captured['payload']['reasoning_budget_tokens'] == 1024
        assert 'response_format' not in captured['payload']


def test_model_json_accepts_a_fence_without_accepting_code():
    assert parse_model_json('Here is the review:\n```json\n{"comments": []}\n```') == {'comments': []}
    with pytest.raises(json.JSONDecodeError):
        parse_model_json('not JSON: __import__("os").system("anything")')


def test_busy_model_request_waits_then_uses_the_same_payload(monkeypatch):
    import io
    from urllib.error import HTTPError
    requests = []
    class Opener:
        def open(self, request, timeout):
            requests.append(request.data)
            if len(requests) == 1:
                raise HTTPError(request.full_url, 409, 'Busy', {}, io.BytesIO(b'{"error":{"message":"Model busy"}}'))
            return io.BytesIO(b'{"choices":[{"message":{"content":"{\\"comments\\": []}"},"finish_reason":"stop"}]}')
    monkeypatch.setattr('engine.review.local_model.build_opener', lambda *args: Opener())
    monkeypatch.setattr('engine.review.local_model.time.sleep', lambda *args: None)
    assert chat({'provider': 'strata', 'base_url': 'http://127.0.0.1:1234', 'model': 'test'},
                [{'role': 'user', 'content': 'Manuscript'}]) == {'comments': []}
    assert len(requests) == 2 and requests[0] == requests[1]


@pytest.mark.parametrize('valid_repair', [True, False])
def test_json_syntax_repair_is_local_and_bounded(monkeypatch, valid_repair):
    import io
    requests = []
    class Opener:
        def open(self, request, timeout):
            requests.append(json.loads(request.data))
            content = '{"comments": []}' if valid_repair and len(requests) == 2 else '{"comments": [}'
            return io.BytesIO(json.dumps({'choices': [{'message': {'content': content}, 'finish_reason': 'stop'}]}).encode())
    monkeypatch.setattr('engine.review.local_model.build_opener', lambda *args: Opener())
    cfg = {'provider': 'strata', 'base_url': 'http://127.0.0.1:1234', 'model': 'test'}
    if valid_repair:
        assert chat(cfg, [{'role': 'user', 'content': 'Manuscript'}]) == {'comments': []}
    else:
        with pytest.raises(LocalReviewError, match='after one syntax-repair attempt'):
            chat(cfg, [{'role': 'user', 'content': 'Manuscript'}])
    assert len(requests) == 2
    assert requests[1]['chat_template_kwargs']['enable_thinking'] is False
    assert cfg['_format_repairs'] == 1


def test_local_model_path_preserves_native_word_gate(tmp_path, monkeypatch):
    from engine.review.selection import select_comments
    from engine.review.persona import compose_all
    from engine.review.endnote import compose_endnote
    from tools.audit_native import audit
    fixture = Path(__file__).parent / 'fixtures/janeen_raw.docx'
    def simulated_model(model, analysis, reviewer, student):
        comments = compose_all(model, select_comments(model, analysis), analysis.essay_type)
        note = compose_endnote(model, analysis, reviewer_name=reviewer, student_name=student, composed=comments)
        return comments, note, {'backend': 'strata', 'model': 'simulated-local-model'}
    monkeypatch.setattr('engine.review.local_model.compose_local', simulated_model)
    result = review_document(fixture, tmp_path, backend='local')
    assert result.validation.ok
    assert not audit(Path(result.output_path))
    assert result.quality['generation']['model'] == 'simulated-local-model'


def test_report_uses_delivered_letter_without_template_regeneration(manuscript, monkeypatch):
    from engine.review.report import build_review
    monkeypatch.setattr('engine.review.report.compose_endnote', lambda *args, **kwargs: pytest.fail('Template letter was regenerated'))
    letter = ['Dear Student,', 'The delivered local-model revision plan.', 'Reviewer']
    report = build_review(manuscript, analyze(manuscript), manuscript.full_text, delivered_endnote=letter)
    assert report['tutor_overview_letter'] == '\n\n'.join(letter)


def test_report_escapes_generated_feedback_html(tmp_path):
    from engine.processor import ReviewProcessor
    processor = ReviewProcessor.__new__(ReviewProcessor)
    target = tmp_path / 'report.html'
    processor.generate_html_report({'executive_summary': '<script>example()</script>',
                                    'top_action_items': ['Discuss <em>the actual wording</em>.'],
                                    'category_scores': {}}, str(target))
    html = target.read_text(encoding='utf8')
    assert '<script>example()</script>' not in html
    assert '&lt;script&gt;example()&lt;/script&gt;' in html
    assert '&lt;em&gt;the actual wording&lt;/em&gt;' in html
