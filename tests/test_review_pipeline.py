"""Regression tests for the deterministic ThinkingStorm review pipeline.

Run with:  python -m unittest discover -s tests -v
"""
from __future__ import annotations

import json
import re
import sys
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine.review.analysis import analyze  # noqa: E402
from engine.review.docmodel import load_document  # noqa: E402
from engine.review.endnote import compose_endnote  # noqa: E402
from engine.review.pipeline import PipelineError, review_document  # noqa: E402
from engine.review.segmenter import split_sentences  # noqa: E402
from engine.review.selection import select_comments  # noqa: E402
from engine.review.style_guard import check_text_block, repeated_ngram_ratio  # noqa: E402
from engine.review.validate import _anchors, _comment_body  # noqa: E402
from engine.review.voice import BANNED_PHRASES, compose_all  # noqa: E402
from tools.extract_review import extract  # noqa: E402

FIXTURES = ROOT / "tests" / "fixtures"
OUT = ROOT / "tests" / "_out"

DOCS = ["janeen_raw.docx", "andrea_raw.docx", "kiran_raw.docx"]


def fixture(name: str) -> Path:
    return FIXTURES / name


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestDocumentModel(unittest.TestCase):
    def test_regions_are_sane(self):
        expected = {
            "janeen_raw.docx": {"style": "mla", "has_title": True, "has_references": True},
            "andrea_raw.docx": {"style": "mla", "has_title": True, "has_references": True},
            "kiran_raw.docx": {"style": "apa", "has_title": True, "has_references": True},
        }
        for name, want in expected.items():
            model = load_document(str(fixture(name)))
            self.assertEqual(model.style, want["style"], name)
            self.assertTrue(any(p.region == "title" for p in model.paragraphs), name)
            self.assertTrue(any(p.region == "references" for p in model.paragraphs), name)
            self.assertGreater(model.word_count, 300, name)

    def test_paragraph_text_roundtrip(self):
        model = load_document(str(fixture("kiran_raw.docx")))
        for p in model.paragraphs:
            self.assertEqual(p.text, "".join(n.text for n in p.nodes))


class TestSegmenter(unittest.TestCase):
    def test_abbreviations_and_citations(self):
        text = (
            "Dr. Smith argues that tuition matters (Burland et al.). "
            "The rate rose 3.5 percent in 2023. "
            "Mr. Jones disagreed; Mrs. Lee agreed."
        )
        sentences = split_sentences(text)
        self.assertEqual(len(sentences), 3)
        for s in sentences:
            self.assertEqual(text[s.start:s.end], s.text)

    def test_quoted_terminal_punctuation(self):
        text = "He called it \u201ca quiet disaster.\u201d The class agreed."
        sentences = split_sentences(text)
        self.assertEqual(len(sentences), 2)


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestRules(unittest.TestCase):
    def test_janeen_findings(self):
        result = analyze(load_document(str(fixture("janeen_raw.docx"))))
        rules = {f.rule_id for f in result.findings}
        self.assertIn("thesis_no_preview", rules)
        self.assertIn("conclusion_summary", rules)
        self.assertIn("mla_header_missing", rules)
        self.assertIn("mla_container_italics", rules)

    def test_kiran_findings(self):
        result = analyze(load_document(str(fixture("kiran_raw.docx"))))
        rules = {f.rule_id for f in result.findings}
        self.assertIn("subject_verb", rules)
        self.assertIn("informal_diction", rules)
        self.assertIn("abstract_intro", rules)
        self.assertIn("apa_date_numeric", rules)

    def test_andrea_findings(self):
        result = analyze(load_document(str(fixture("andrea_raw.docx"))))
        rules = {f.rule_id for f in result.findings}
        self.assertIn("mla_date_numeric", rules)
        self.assertIn("mla_heading_order", rules)
        self.assertIn("thesis_rhetorical_missing", rules)
        self.assertIn("fragment", rules)

    def test_no_front_matter_transition_noise(self):
        """A heading line must never be diagnosed as an abrupt transition."""
        for name in DOCS:
            model = load_document(str(fixture(name)))
            result = analyze(model)
            for f in result.findings:
                if f.rule_id == "transition_abrupt":
                    self.assertEqual(f.paragraph.region, "body", f"{name}:{f.quote}")

    def test_rhetorical_essay_does_not_get_apa_rules(self):
        model = load_document(str(fixture("andrea_raw.docx")))
        result = analyze(model)
        self.assertFalse([f for f in result.findings if f.category == "APA"])


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestVoice(unittest.TestCase):
    def test_comments_avoid_banned_phrases(self):
        for name in DOCS:
            model = load_document(str(fixture(name)))
            result = analyze(model)
            composed = compose_all(model, select_comments(model, result), result.essay_type)
            for _item, body in composed:
                low = body.lower()
                for phrase in BANNED_PHRASES:
                    self.assertNotIn(phrase, low, f"{name}: {phrase} in {body!r}")

    def test_composition_is_deterministic(self):
        model = load_document(str(fixture("kiran_raw.docx")))
        result = analyze(model)
        first = [b for _i, b in compose_all(model, select_comments(model, result), result.essay_type)]
        second = [b for _i, b in compose_all(model, select_comments(model, result), result.essay_type)]
        self.assertEqual(first, second)

    def test_opening_phrase_is_not_repeated(self):
        for name in DOCS:
            model = load_document(str(fixture(name)))
            result = analyze(model)
            composed = compose_all(model, select_comments(model, result), result.essay_type)
            openers = [" ".join(b.split()[:4]).lower() for _i, b in composed]
            self.assertGreaterEqual(len(set(openers)), max(3, int(len(openers) * 0.7)), name)


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestEndnote(unittest.TestCase):
    def test_domains_and_length(self):
        for name in DOCS:
            model = load_document(str(fixture(name)))
            result = analyze(model)
            endnote = compose_endnote(model, result)
            self.assertEqual(len(endnote.domains), 3, name)
            self.assertTrue(endnote.paragraphs[0].startswith("Dear "), name)
            self.assertEqual(endnote.paragraphs[-1], "Raunak Banga")
            self.assertTrue(150 <= endnote.word_count <= 320, f"{name}: {endnote.word_count}")
            headings = [p for p in endnote.paragraphs if p.endswith(":")]
            self.assertGreaterEqual(len(headings), 3, name)

    def test_letters_differ_between_documents(self):
        texts = []
        for name in DOCS:
            model = load_document(str(fixture(name)))
            result = analyze(model)
            texts.append(" ".join(compose_endnote(model, result).paragraphs))
        self.assertLess(repeated_ngram_ratio(texts, n=6), 0.35)


class TestStyleGuard(unittest.TestCase):
    def test_human_comments_pass(self):
        exemplar = ROOT / "style_corpus" / "exemplars.json"
        if not exemplar.exists():
            self.skipTest("style corpus missing")
        data = json.loads(exemplar.read_text(encoding="utf-8"))
        by_doc: dict[str, list[str]] = {}
        for comment in data["comments"]:
            by_doc.setdefault(comment["doc"], []).append(comment["body"])
        for name, bodies in by_doc.items():
            report = check_text_block(bodies, name)
            self.assertTrue(report.ok, f"{name}: {report.errors}")

    def test_repetitive_text_is_flagged(self):
        bodies = ["The paragraph presents valuable evidence but the analysis is too summarizing."] * 6
        report = check_text_block(bodies, "synthetic")
        self.assertFalse(report.ok)


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestPipeline(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        OUT.mkdir(parents=True, exist_ok=True)
        cls.results = {}
        for name in DOCS:
            cls.results[name] = review_document(
                str(fixture(name)), output_dir=str(OUT), overwrite=True
            )

    def test_deliverables_validate(self):
        for name, result in self.results.items():
            self.assertTrue(result.validation.ok, f"{name}: {result.validation.errors}")
            self.assertTrue(Path(result.output_path).exists())
            self.assertTrue(Path(result.report_path).exists())

    def test_comment_budget(self):
        for name, result in self.results.items():
            self.assertGreaterEqual(len(result.comments), 5, name)
            self.assertLessEqual(len(result.comments), 22, name)

    def test_quotes_are_exact_and_unique(self):
        for name, result in self.results.items():
            data = extract(Path(result.output_path))
            quotes = [c["quote"] for c in data["comments"]]
            self.assertEqual(len(quotes), len(set(quotes)), name)
            for c in data["comments"]:
                self.assertTrue(c["quote"].strip(), f"{name}: empty anchor")

    def test_endnote_typography(self):
        for name, result in self.results.items():
            with zipfile.ZipFile(result.output_path) as z:
                xml = z.read("word/endnotes.xml").decode("utf-8")
            self.assertIn("Times New Roman", xml)
            self.assertIn('w:color w:val="FF0000"', xml)
            self.assertIn('w:sz w:val="24"', xml)
            self.assertIn('w:line="480"', xml)

    def test_links_are_allowlisted(self):
        with zipfile.ZipFile(self.results["janeen_raw.docx"].output_path) as z:
            rels = z.read("word/_rels/comments.xml.rels").decode("utf-8")
        targets = re.findall(r'Target="([^"]+)"', rels)
        self.assertTrue(targets)
        for target in targets:
            self.assertTrue(target.startswith("https://knowledge.thinkingstorm.com/"), target)

    def test_second_run_creates_a_new_file(self):
        again = review_document(str(fixture("janeen_raw.docx")), output_dir=str(OUT))
        self.assertNotEqual(again.output_path, self.results["janeen_raw.docx"].output_path)

    def test_already_commented_file_is_rejected(self):
        original = Path(r"C:/Users/shashank/Downloads/janeen_REV_BY_TS_RaunakB_243347_1.docx")
        if not original.exists():
            self.skipTest("original reviewed file not available")
        with self.assertRaises(PipelineError):
            review_document(str(original), output_dir=str(OUT))


@unittest.skipUnless(FIXTURES.exists(), "fixtures not built")
class TestStatementGenre(unittest.TestCase):
    def test_personal_statement_path(self):
        import docx

        OUT.mkdir(parents=True, exist_ok=True)
        source = ROOT / "samples" / "sample_sop.txt"
        if not source.exists():
            self.skipTest("sample statement missing")
        doc = docx.Document()
        for para in [p.strip() for p in source.read_text(encoding="utf-8").split("\n\n") if p.strip()]:
            doc.add_paragraph(para)
        draft = OUT / "statement_raw.docx"
        doc.save(str(draft))

        result = review_document(str(draft), output_dir=str(OUT), overwrite=True)
        self.assertEqual(result.analysis.essay_type, "statement")
        self.assertTrue(result.validation.ok, result.validation.errors)
        self.assertGreaterEqual(len(result.comments), 3)
        rules = {c["rule"] for c in result.comments}
        self.assertIn("statement_cliche_opener", rules)


class TestExistingComments(unittest.TestCase):
    def test_append_to_reviewed_file(self):
        reviewed = Path(r"C:/Users/shashank/Downloads/janeen_REV_BY_TS_RaunakB_243347_1.docx")
        if not reviewed.exists():
            self.skipTest("reviewed exemplar not available")
        OUT.mkdir(parents=True, exist_ok=True)
        result = review_document(str(reviewed), output_dir=str(OUT),
                                 allow_existing_comments=True, overwrite=True)
        self.assertTrue(result.validation.ok, result.validation.errors)
        self.assertGreater(result.validation.metrics["pre_existing_comments"], 0)


class TestVoiceTemplates(unittest.TestCase):
    LABELS = [
        "THESIS STATEMENT", "INTRODUCTION", "ORGANISATION", "TOPIC SENTENCE",
        "TRANSITION", "ANALYSIS OF EVIDENCE", "STATISTICS", "ANALYSIS",
        "COUNTER-ARGUMENT", "CONCLUSION", "ABSTRACT", "MLA", "APA",
        "PUNCTUATION", "FRAGMENTS", "SUBJECT-VERB AGREEMENT", "INFORMAL WRITING",
        "SENTENCE CLARITY",
    ]

    def test_templates_do_not_start_with_a_category_label(self):
        from engine.review.voice import TEMPLATES

        for rule, variants in TEMPLATES.items():
            for lead, _advice in variants:
                for label in self.LABELS:
                    self.assertFalse(
                        lead.upper().startswith(label + " "),
                        f"{rule} would duplicate its label: {lead[:60]!r}",
                    )


class TestTables(unittest.TestCase):
    def test_table_paragraphs_are_not_treated_as_prose(self):
        import docx

        OUT.mkdir(parents=True, exist_ok=True)
        doc = docx.Document()
        doc.add_paragraph("Marcus Chen")
        doc.add_paragraph("Professor Alvarez")
        doc.add_paragraph("ENG 210")
        doc.add_paragraph("12 October 2026")
        doc.add_paragraph("The Case for Later Start Times")
        for para in (
            "Schools across the country begin classes before eight in the morning, and students pay for it. "
            "Research on adolescent sleep shows that the teenage circadian rhythm shifts later during puberty.",
            "The evidence on sleep and learning is consistent, and districts that moved the first bell have "
            "seen attendance rise. The researchers found that students gained nearly an hour of sleep each "
            "night, and the gains appeared across every demographic group in the sample. Attendance gains "
            "held steady through the second year of the study.",
            "Later start times are not a complete solution, but they are the cheapest and best-supported "
            "change available. Schools should move the first bell to 8:30.",
        ):
            doc.add_paragraph(para)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "District"
        table.cell(0, 1).text = "8:45"
        path = OUT / "table_case.docx"
        doc.save(str(path))

        result = review_document(str(path), output_dir=str(OUT), overwrite=True)
        self.assertTrue(result.validation.ok, result.validation.errors)
        for comment in result.comments:
            self.assertNotIn("8:45", comment["quote"])
            self.assertNotIn("District", comment["quote"])


class TestNoModelDependency(unittest.TestCase):
    def test_pipeline_modules_do_not_import_network_or_llm_clients(self):
        package = ROOT / "engine" / "review"
        forbidden = ("import openai", "import anthropic", "import requests",
                     "import httpx", "from openai", "from anthropic", "from requests")
        for path in package.glob("*.py"):
            text = path.read_text(encoding="utf-8")
            for token in forbidden:
                self.assertNotIn(token, text, f"{path.name} imports {token}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
