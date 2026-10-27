"""Regressions for the sparse-review failure reported in Studio."""
from pathlib import Path
import tempfile
from zipfile import ZipFile

import docx
from docx.oxml import OxmlElement

from engine.review.analysis import AnalysisResult, Finding, PUNCT, analyze, build_blocks
from engine.review.coverage import page_map, assess
from engine.review.docmodel import load_document
from engine.review.endnote import Endnote
from engine.review.selection import select_comments
from engine.review.source_material import remove_previous_letter


def save_document(directory, paragraphs):
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    path = Path(directory) / 'draft.docx'
    document.save(path)
    return path


def test_substantial_unpunctuated_paragraphs_remain_separate():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['A student discusses the use of technology in daily life ' * 8,
                                         'Another point concerns the effects on education and employment. ' * 8])
        model = load_document(str(path))
        assert len(build_blocks(model)) == 2
        assert 'paragraph_terminal_punctuation' in {f.rule_id for f in analyze(model).findings}


def test_short_acronym_anchor_survives_analysis():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['Artificial intelligence is the subject of this discussion. ' * 7 +
                                         'The use of ai requires careful consideration.'])
        findings = analyze(load_document(str(path))).findings
        assert any(f.rule_id == 'acronym_case' and f.quote == 'ai' for f in findings)


def test_period_spacing_has_local_findings_even_with_proofreading_summary():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['The tools can assist with problem solving . Being able to check results matters. '
                                         'The tool is wrong sometimes.If the designers change it, users should check again. ' * 3])
        findings = analyze(load_document(str(path))).findings
        assert any(f.rule_id == 'space_before_punct' for f in findings)
        assert any(f.rule_id == 'missing_space_after_punct' for f in findings)


def test_bibliography_sources_are_not_hidden_in_front_matter_or_body():
    with tempfile.TemporaryDirectory() as directory:
        title = 'Annotated Bibliography: How does access to legal representation affect the outcomes of immigration court cases and applications?'
        source = 'Smith, A. (2025). A study of access. Example Journal, 4, 1-9. https://example.test/study'
        path = save_document(directory, [title, 'Avery Smith', source,
                                         'The authors describe the research and explain their findings in this study. ' * 6])
        model = load_document(str(path))
        assert model.paragraphs[0].region == 'title'
        assert model.paragraphs[2].region == 'references'
        assert source not in ' '.join(p.text for p in model.body_paragraphs)
        assert analyze(model).essay_type == 'bibliography'


def test_old_letter_is_excluded_and_cleanup_preserves_other_package_parts():
    with tempfile.TemporaryDirectory() as directory:
        manuscript = 'The student discusses a specific experience in this passage. ' * 10
        path = save_document(directory, [manuscript, 'Hi Avery,',
                                         'You have done a good job describing your experience.',
                                         'There are a few areas to develop in your next draft.', 'Best wishes,', 'Reviewer'])
        original = path.read_bytes()
        model = load_document(str(path))
        assert model.full_text == manuscript
        assert model.word_count == len(manuscript.split())
        target = Path(directory) / 'clean.docx'
        assert remove_previous_letter(path, target)['removed']
        assert path.read_bytes() == original
        assert load_document(str(target)).full_text == manuscript
        with ZipFile(path) as before, ZipFile(target) as after:
            assert before.namelist() == after.namelist()
            assert all(before.read(n) == after.read(n) for n in before.namelist() if n != 'word/document.xml')


def test_student_letter_is_not_removed():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['Dear Professor,', 'I am sharing my next draft for your consideration. '
                                         'The discussion concerns the experience described in class. ' * 5,
                                         'Best wishes,', 'Avery'])
        result = remove_previous_letter(path, Path(directory) / 'clean.docx')
        assert not result['removed']
        assert 'Dear Professor' in load_document(str(path)).full_text


def test_cached_pagination_is_labelled_and_anchors_follow_breaks():
    with tempfile.TemporaryDirectory() as directory:
        document = docx.Document()
        paragraph = document.add_paragraph('First passage ends here. ' * 5)
        offset = len(paragraph.text)
        paragraph.add_run()._r.append(OxmlElement('w:lastRenderedPageBreak'))
        paragraph.add_run('Second passage starts here. ' * 5)
        path = Path(directory) / 'draft.docx'
        document.save(path)
        model = load_document(str(path))
        pages = page_map(model)
        first = Finding('double_period', PUNCT, model.paragraphs[0], 0, 10, {}, 60)
        second = Finding('double_period', PUNCT, model.paragraphs[0], offset, offset + 10, {}, 60)
        assert pages.pages == [1, 2]
        assert pages.page_for(first) == 1
        assert pages.page_for(second) == 2
        quality = assess(model, [], Endnote([], [], 0))
        assert quality['page_count_is_estimate']
        assert quality['undercovered_pages'] == [1, 2]
        assert quality['needs_review']


def test_sparse_page_is_flagged_without_inventing_findings():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['A student discusses an experience and explains its relevance. ' * 15])
        model = load_document(str(path))
        para = model.body_paragraphs[0]
        finding = Finding('double_period', PUNCT, para, 0, 10, {}, 60)
        selected = select_comments(model, AnalysisResult([finding], 'reflection', [], {}, []), comprehensive=True)
        assert len(selected) == 1
        quality = assess(model, [(selected[0], 'Remove the extra period here.')], Endnote([], [], 0))
        assert quality['needs_review']
        assert quality['comments_by_page'] == [{'page': 1, 'comments': 1}]


def test_statistical_interpretation_and_analysis_require_development():
    from engine.review.analysis import ANALYSIS, STATISTICS
    from engine.review.selection import SelectedComment
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['The student reports a result and explains the discussion. ' * 8])
        model = load_document(str(path))
        for category in (ANALYSIS, STATISTICS):
            finding = Finding('assessment_interpretation', category, model.body_paragraphs[0], 0, 10, {}, 70)
            item = SelectedComment(finding, category, True)
            quality = assess(model, [(item, 'Explain this result.')], Endnote([], [], 0))
            assert quality['developmental_comments'] == 1
            assert quality['shallow_developmental_rules'] == ['assessment_interpretation']


def test_mid_sentence_extraction_is_not_diagnosed_as_missing_punctuation():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['The discussion considers the cost of higher education. ' * 8 + 'Students need affordable',
                                         'education to complete their degrees and continue their studies.'])
        model = load_document(str(path))
        assert len(build_blocks(model)) == 1
        assert not any(f.rule_id == 'paragraph_terminal_punctuation' for f in analyze(model).findings)


def test_https_location_is_not_mistaken_for_a_container_title():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['The student discusses a source and explains the argument. ' * 8,
                                         'Works Cited',
                                         'Smith, Avery. “A useful study.” Example Magazine, 2025, https://www.example.test/study.'])
        model = load_document(str(path), academic_style='mla')
        assert 'mla_raw_url_container' not in {f.rule_id for f in analyze(model).findings}


def test_authorless_entry_does_not_treat_website_as_article_title():
    with tempfile.TemporaryDirectory() as directory:
        path = save_document(directory, ['The student discusses solar energy and evaluates a source. ' * 8,
                                         'Works Cited',
                                         'Solar energy and its applications. (n.d.). Engineering | Example University. https://example.test/solar',
                                         'Smith, A. (2025). A study of solar energy. Science Review. https://example.test/study'])
        findings = analyze(load_document(str(path), academic_style='mla')).findings
        titles = [f.quote for f in findings if f.rule_id == 'mla_article_title_format']
        assert 'Engineering | Example University.' not in titles
        assert titles == ['A study of solar energy.']


def test_watcher_does_not_label_sparse_review_as_complete():
    from engine.review.watch import process_once
    with tempfile.TemporaryDirectory() as directory:
        inbox = Path(directory) / 'inbox'
        inbox.mkdir()
        save_document(inbox, ['A student reflects on a classroom experience and describes the lesson learned. ' * 8])
        results = process_once(inbox, Path(directory) / 'outbox')
        assert len(results) == 1
        assert results[0]['status'] == 'needs_review'
        assert results[0]['review_quality']['needs_review']


def test_studio_propagates_quality_warning_in_response_and_saved_status():
    from unittest.mock import patch
    from portal import server
    def one_verified_issue(model):
        para = next(p for p in model.body_paragraphs if 'ends here..' in p.text)
        start = para.text.index('ends here..')
        finding = Finding('double_period', PUNCT, para, start, start + len('ends here..'), {}, 61)
        return AnalysisResult([finding], 'reflection', [], {}, [])
    with tempfile.TemporaryDirectory() as directory, patch.object(server, 'OUTPUT_DIR', Path(directory)), \
            patch.object(server, 'get_db') as database, \
            patch('engine.review.pipeline.analyze', side_effect=one_verified_issue):
        task = server.execute_thinkingstorm_pipeline(
            student_name='Avery', reviewer_name='Reviewer', reviewer_initials='R', academic_style='mla',
            text='A student reflects on a classroom experience and describes the lesson learned. ' * 8 + 'The account ends here..',
            custom_comments_text='', custom_endnote_text='', file=None)
        import asyncio
        payload = asyncio.run(task)
        assert payload['status'] == 'needs_review'
        assert payload['review_quality']['needs_review']
        assert 'Coverage needs review' in payload['warning']
        parameters = database.return_value.__enter__.return_value.execute.call_args.args[1]
        assert 'needs_review' in parameters
